"""
ingest.py — Parse documents and build the local vector index.

Supported formats: PDF, DOCX, Markdown, plain text (.txt / .rst)

Usage:
    python ingest.py
    python ingest.py path/to/specific/file.pdf   # re-index a single file
"""

import os
import sys
import hashlib

import chromadb
from sentence_transformers import SentenceTransformer

import config


# ── Text extraction ────────────────────────────────────────────────────────────

def _extract_pdf(path: str) -> str:
    import fitz  # pymupdf
    doc = fitz.open(path)
    return "\n".join(page.get_text() for page in doc)


def _extract_docx(path: str) -> str:
    from docx import Document
    doc = Document(path)
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    # Also pull text from tables
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
    return "\n".join(parts)


def _extract_text(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as fh:
        return fh.read()


def extract_text(path: str) -> str | None:
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            return _extract_pdf(path)
        elif ext == ".docx":
            return _extract_docx(path)
        elif ext in (".md", ".txt", ".rst"):
            return _extract_text(path)
    except Exception as exc:
        print(f"    [error] Could not read {os.path.basename(path)}: {exc}")
    return None


# ── Chunking ───────────────────────────────────────────────────────────────────

def chunk_text(text: str) -> list[str]:
    """Split text into overlapping fixed-size chunks."""
    chunks: list[str] = []
    start = 0
    size = config.CHUNK_SIZE
    overlap = config.CHUNK_OVERLAP
    while start < len(text):
        chunk = text[start : start + size].strip()
        if chunk:
            chunks.append(chunk)
        start += size - overlap
    return chunks


# ── Helpers ────────────────────────────────────────────────────────────────────

def _file_id(path: str) -> str:
    """Stable ID for a file path (used as a ChromaDB filter key)."""
    return hashlib.md5(os.path.abspath(path).encode()).hexdigest()


def _collect_files(root: str) -> list[str]:
    supported = {".pdf", ".docx", ".md", ".txt", ".rst"}
    found: list[str] = []
    for dirpath, _, filenames in os.walk(root):
        for name in filenames:
            if os.path.splitext(name)[1].lower() in supported:
                found.append(os.path.join(dirpath, name))
    return found


# ── Main ───────────────────────────────────────────────────────────────────────

def ingest(target_files: list[str] | None = None) -> None:
    # Ensure docs directory exists
    if not os.path.isdir(config.DOCS_DIR):
        os.makedirs(config.DOCS_DIR)
        print(f"Created '{config.DOCS_DIR}/' — add your documents there and run again.")
        return

    files = target_files if target_files else _collect_files(config.DOCS_DIR)
    if not files:
        print(f"No supported documents found in '{config.DOCS_DIR}/'. Nothing to index.")
        return

    print(f"Loading embedding model '{config.EMBEDDING_MODEL}'...")
    model = SentenceTransformer(config.EMBEDDING_MODEL)

    client = chromadb.PersistentClient(path=config.DB_DIR)
    collection = client.get_or_create_collection(
        name="docs",
        metadata={"hnsw:space": "cosine"},
    )

    print(f"Found {len(files)} file(s). Indexing...\n")
    total_chunks = 0

    for filepath in files:
        basename = os.path.basename(filepath)
        print(f"  {basename}")

        text = extract_text(filepath)
        if not text or not text.strip():
            print("    → skipped (no extractable text)")
            continue

        chunks = chunk_text(text)
        if not chunks:
            print("    → skipped (empty after chunking)")
            continue

        fid = _file_id(filepath)

        # Remove any previously indexed chunks for this file so re-runs are safe
        existing = collection.get(where={"source_id": fid})
        if existing["ids"]:
            collection.delete(ids=existing["ids"])
            print(f"    → removed {len(existing['ids'])} old chunk(s)")

        embeddings = model.encode(chunks, show_progress_bar=False).tolist()
        ids = [f"{fid}_{i}" for i in range(len(chunks))]
        metadatas = [
            {"source": basename, "source_id": fid, "chunk_index": i}
            for i in range(len(chunks))
        ]

        collection.add(ids=ids, embeddings=embeddings, documents=chunks, metadatas=metadatas)
        total_chunks += len(chunks)
        print(f"    → {len(chunks)} chunks indexed")

    print(f"\nDone. {total_chunks} chunk(s) stored in '{config.DB_DIR}/'.")
    print(f"Total chunks in index: {collection.count()}")


if __name__ == "__main__":
    explicit = sys.argv[1:]  # optional: pass specific file paths as arguments
    ingest(explicit if explicit else None)
